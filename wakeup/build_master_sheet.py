"""Genera WAKE_UP_Master.xlsx: el tablero de producción de WAKE UP!
(contexto, mapa de la experiencia, guiones, ejercicios, transiciones FM,
producción, frases, investigación e ideas). Se sube a Google Drive como
Google Sheet; a partir de ahí la hoja de Drive es la fuente de verdad."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

OUT = "wakeup/WAKE_UP_Master.xlsx"

INK = "111111"
PAPER = "FFFFFF"
HEAD_FILL = PatternFill("solid", fgColor=INK)
HEAD_FONT = Font(name="Arial", bold=True, color=PAPER, size=10)
BODY_FONT = Font(name="Arial", size=10)
TITLE_FONT = Font(name="Arial", bold=True, size=16)
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

FORMAT_COLORS = {
    "TRANSICIÓN FM": "FFD8A8",
    "AUDIO CINEMATOGRÁFICO": "D0BFFF",
    "RAW VIDEO": "FFF3A3",
    "RAW AUDIO / PODCAST": "FFE08A",
    "EXPERIENCIA GUIADA": "A5D8FF",
    "EJERCICIO CREATIVO": "B2F2BB",
    "LIBRO + RAW VIDEO": "F1E3C8",
    "VIDEO CINEMATOGRÁFICO": "FFC9C9",
}

ESTADOS = ["IDEA", "BORRADOR", "EN REVISIÓN", "APROBADO", "GRABADO", "EDITADO", "FINAL"]
PROD = ["PENDIENTE", "EN PROCESO", "LISTO", "N/A"]
ESTADO_COLORS = {
    "IDEA": "EEEEEE", "BORRADOR": "FFF3BF", "EN REVISIÓN": "FFE8CC",
    "APROBADO": "D3F9D8", "GRABADO": "C5F6FA", "EDITADO": "D0EBFF", "FINAL": "8CE99A",
    "PENDIENTE": "EEEEEE", "EN PROCESO": "FFE8CC", "LISTO": "8CE99A",
}

wb = Workbook()


def sheet(title, headers, widths, rows, estado_cols=(), prod_cols=(), fmt_col=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.alignment, c.border = HEAD_FILL, HEAD_FONT, Alignment(wrap_text=True, vertical="center"), BORDER
    ws.row_dimensions[1].height = 30
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.font, c.alignment, c.border = BODY_FONT, WRAP, BORDER
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "B2" if len(headers) > 4 else "A2"
    ws.auto_filter.ref = ws.dimensions
    last = max(ws.max_row, 2) + 200

    def dropdown(cols, values):
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        for col in cols:
            L = ws.cell(1, col).column_letter
            dv.add(f"{L}2:{L}{last}")
            for v, color in ESTADO_COLORS.items():
                if v in values:
                    ws.conditional_formatting.add(
                        f"{L}2:{L}{last}",
                        FormulaRule(formula=[f'{L}2="{v}"'], fill=PatternFill("solid", fgColor=color)))

    if estado_cols:
        dropdown(estado_cols, ESTADOS)
    if prod_cols:
        dropdown(prod_cols, PROD)
    if fmt_col:
        L = ws.cell(1, fmt_col).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(FORMAT_COLORS) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{L}2:{L}{last}")
        for v, color in FORMAT_COLORS.items():
            ws.conditional_formatting.add(
                f"{L}2:{L}{last}",
                FormulaRule(formula=[f'{L}2="{v}"'], fill=PatternFill("solid", fgColor=color)))
    return ws


# ---------------------------------------------------------------- 00 LÉEME
ws = wb.active
ws.title = "00 LÉEME"
ws["A1"] = "WAKE UP! — TABLERO MAESTRO"
ws["A1"].font = TITLE_FONT
ws["A2"] = "LA CREATIVIDAD NO SE APRENDE. SE RECUERDA.  ·  DESPIERTA. RECUERDA. JUEGA. CREA."
ws["A2"].font = Font(name="Arial", italic=True, size=11)
guide = [
    ("PESTAÑA", "PARA QUÉ SIRVE"),
    ("01 CONTEXTO", "La biblia del proyecto: qué es WAKE UP!, qué no es, filosofía, tono y regla principal. Léela antes de escribir cualquier pieza."),
    ("02 MAPA", "La receta completa, pieza por pieza y en orden. Cada fila es un ingrediente: formato, objetivo, qué siente y qué hace la persona, y cómo prepara la siguiente pieza."),
    ("03 GUIONES", "Los guiones, divididos en bloques (una fila por bloque). Columnas: VOZ, SONIDO, VISUAL y DIRECCIÓN. Filtra por ID para ver una sola pieza. Ya hay borradores de INTRO, 01, 02, 03, 05 y 06."),
    ("04 EJERCICIOS", "Banco de misiones creativas, con reglas, tiempo, materiales y resultado. La columna PIEZA dice dónde vive cada una."),
    ("05 TRANSICIONES FM", "Las transmisiones de radio que cosen el recorrido: frecuencia, mensaje y diseño sonoro."),
    ("06 PRODUCCIÓN", "El checklist por pieza: guion → grabación → edición → sonido → revisión → final."),
    ("07 FRASES", "Banco de frases, títulos y copy. Si una frase se usa en una pieza, se anota ahí."),
    ("08 INVESTIGACIÓN", "Afirmaciones científicas o históricas que queramos usar. Aquí se separa la FILOSOFÍA / METÁFORA de WAKE UP! de lo que tiene EVIDENCIA. Nada entra a un guion como dato sin pasar por aquí."),
    ("09 IDEAS", "El parqueadero: ideas sueltas, referencias, cosas para PLAYGROUND. Se anota sin juzgar."),
    ("", ""),
    ("CÓMO USAR", "Los desplegables de ESTADO y FORMATO se colorean solos. Los tiempos de los guiones son aproximados. En los guiones, [PAUSA 5s] marca silencios, (entre paréntesis) va una indicación de interpretación y lo que está en MAYÚSCULAS es una palabra con peso."),
    ("REGLA DE ORO", "Ante cada pieza, pregúntate: ¿solo está enseñando creatividad o está haciendo que la persona la experimente? Siempre la segunda. La persona debe salir habiendo creado algo."),
]
for i, (a, b) in enumerate(guide, start=4):
    ws.cell(i, 1, a).font = Font(name="Arial", bold=True, size=10, color=PAPER if i == 4 else INK)
    ws.cell(i, 2, b).font = Font(name="Arial", size=10, bold=(i == 4), color=PAPER if i == 4 else INK)
    ws.cell(i, 2).alignment = WRAP
    if i == 4:
        ws.cell(i, 1).fill = ws.cell(i, 2).fill = HEAD_FILL
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 110

# ---------------------------------------------------------------- 01 CONTEXTO
contexto = [
    ("QUÉ ES", "Un programa de DESPERTAR CREATIVO. No es un curso convencional ni una colección de clases teóricas: es una experiencia audiovisual, introspectiva, práctica y narrativa que acompaña a la persona a reconectar con la capacidad de crear que ya existe dentro de ella."),
    ("PREMISA", "La creatividad no es algo que tengas que aprender desde cero. Es algo que tienes que recordar. Si algo necesita despertarse, significa que alguna vez estuvo despierto, y uno de los momentos en que más despierto estuvo fue la infancia."),
    ("QUÉ LA DURMIÓ", "El miedo, los prejuicios, la comparación, las responsabilidades, las expectativas externas, la búsqueda de aprobación y el exceso de pensamiento."),
    ("EL NIÑO INTERIOR", "Explora sin saber qué busca, juega sin necesitar productividad, crea sin pensar si será rentable, pregunta, imagina mundos, se equivoca sin crisis y hace cosas porque disfruta hacerlas. No se trata de volver a ser niño, sino de recuperar curiosidad, juego, imaginación, libertad y asombro con la experiencia y la conciencia de un adulto."),
    ("TODOS SOMOS CREADORES", "Crear una empresa, una canción, una solución, una conversación, una película, un producto, una relación, un movimiento, una nueva versión de ti. No haces cosas creativas: ERES creador. Principio maestro: «Crear no es hacer algo… crear es ser»."),
    ("OBJETIVO", "No decirle a nadie qué crear, sino devolverle el estado en el que quiere explorar → sentir → imaginar → jugar → descubrir → crear → materializar. Empieza con «No sé si soy creativo» y termina con «Crear siempre estuvo dentro de mí». WAKE UP! es la semilla."),
    ("PILAR · SE RECUERDA", "«La creatividad es algo que se recuerda, no algo que se aprende». No se trata a la persona como alguien incompleto."),
    ("PILAR · POR AMOR AL ARTE", "Crear sin que el primer pensamiento sea el dinero, la fama, el estatus, los likes o la productividad. Primero el amor, la curiosidad, el juego y la exploración; los resultados externos vienen después."),
    ("PILAR · FLUIR", "No se elimina el pensamiento racional: hay momentos para analizar y momentos para dejar aparecer las ideas. Se crean espacios para divagar, conectar, observar, escribir, imaginar y jugar."),
    ("PILAR · IDEA → ACCIÓN", "Tener ideas no es suficiente: una idea se agradece haciéndola realidad. IDEA → ACCIÓN → CREACIÓN, aunque sea imperfecta."),
    ("PILAR · CREER → CREAR → CRECER", "Primero recuperar la posibilidad. Después crear. Después dejar que lo creado nos transforme."),
    ("NO ES UN CURSO", "Debe sentirse como entrar a una película, una transmisión o un videojuego mental. La persona escucha, observa, piensa, escribe, respira, imagina, recuerda, experimenta y crea. Por eso los formatos cambian todo el tiempo."),
    ("FORMATOS", "RAW VIDEO (Camilo a cámara, íntimo e imperfecto) · RAW AUDIO / PODCAST · AUDIO CINEMATOGRÁFICO (voz, música, ambientes, silencio) · VIDEO CINEMATOGRÁFICO (simbólico) · EXPERIENCIA GUIADA (introspección y visualización; nunca se presenta como tratamiento médico o psicológico) · EJERCICIO CREATIVO (en el mundo real) · TRANSICIONES FM (interferencias, frecuencias, cambios de canal) · LIBRO / TEXTOS / CARTAS."),
    ("ES UNA RECETA", "Es un recorrido en secuencia, no una biblioteca para ver al azar. Cada pieza prepara la siguiente, y si falta una parte debe sentirse que falta un ingrediente."),
    ("EL FINAL", "EL NIÑO ESTÁ DESPIERTO → ¿Y AHORA QUÉ? → JUGAR → PLAYGROUND. WAKE UP! despierta y PLAYGROUND es el lugar para jugar, explorar y crear."),
    ("TONO · SÍ", "Cinematográfico, crudo, humano, profundo, experimental, espiritual sin obligar a nadie a serlo, infantil sin ser infantilizado, filosófico pero comprensible, poético pero práctico, emocional pero no artificial, misterioso pero claro. Una película + un álbum + un libro + una conversación + un juego."),
    ("TONO · NO", "Nada de «curso online», gurú, «10 hacks para ser creativo», productividad extrema ni motivación barata."),
    ("UNIVERSO / PALABRAS", "despertar · sueño · frecuencias · señales · transmisiones · interferencia · radio · recuerdos · infancia · juego · imaginación · semillas · ideas · descargas · flow · curiosidad · exploración · creación · universo · conciencia · arte · amor · playground"),
    ("CIENCIA VS. METÁFORA", "Las ideas espirituales y metafóricas son parte de la narrativa y no se eliminan, pero nunca se presentan como hechos científicos si no hay evidencia. Ver la pestaña 08 INVESTIGACIÓN."),
    ("REGLA PRINCIPAL", "¿Esto solo enseña creatividad o hace que la persona la EXPERIMENTE? Siempre la segunda. La persona debe salir del programa habiendo creado algo."),
    ("FRASE CENTRAL", "WAKE UP! YOU DON'T LEARN CREATIVITY. YOU REMEMBER IT.  /  WAKE UP! LA CREATIVIDAD NO SE APRENDE. SE RECUERDA."),
]
sheet("01 CONTEXTO", ["TEMA", "CONTENIDO"], [28, 120], contexto)

# ---------------------------------------------------------------- 02 MAPA
mapa_h = ["ID", "ORDEN", "FASE", "FORMATO", "TÍTULO PROVISIONAL", "PREGUNTA / OBJETIVO",
          "QUÉ SIENTE LA PERSONA", "QUÉ HACE LA PERSONA", "DURACIÓN EST.",
          "INGREDIENTE (por qué no puede faltar)", "PREPARA →", "ESTADO", "LINK GUION / ARCHIVO", "NOTAS"]
mapa = [
    ("INTRO", 0, "0 · UMBRAL", "TRANSICIÓN FM", "SEÑAL",
     "Entrar al universo WAKE UP!. Dejar claro que esto no es una clase, sino una transmisión.",
     "Curiosidad, misterio, la sensación de haber encontrado algo por accidente.",
     "Se pone audífonos, busca un lugar tranquilo y deja lista la libreta.",
     "1–2 min", "Marca el pacto: aquí se escucha, se siente y se hace, no solo se aprende.", "01", "BORRADOR", "", ""),
    ("01", 1, "1 · DESPIERTA", "AUDIO CINEMATOGRÁFICO", "ALGO SE APAGÓ",
     "¿Qué es la creatividad? Plantear la posibilidad de que algo dentro de nosotros se haya apagado.",
     "Nostalgia, reconocimiento («eso me pasó a mí»), una pequeña herida que se abre con cariño.",
     "Escucha con los ojos cerrados y recuerda la última vez que perdió la noción del tiempo creando.",
     "4–6 min", "Instala la premisa: la creatividad se recuerda, no se aprende.", "02", "BORRADOR", "", ""),
    ("02", 2, "1 · DESPIERTA", "RAW VIDEO", "EL DESPERTAR CREATIVO",
     "Poner cara y verdad: el despertar de Camilo, con la realidad, la creatividad, la espiritualidad y el cambio.",
     "Confianza. «Esto no me lo está vendiendo un gurú; lo vivió alguien como yo».",
     "Escucha y conecta. Al final recibe la invitación a la primera experiencia.",
     "6–10 min", "Vuelve humano el concepto y le da permiso a la persona para creer que le puede pasar.", "03", "BORRADOR", "", ""),
    ("03", 3, "1 · DESPIERTA", "EXPERIENCIA GUIADA", "LLAMAR A LO QUE DUERME",
     "Primera introspección profunda: empezar a llamar a lo que está dormido (primer encuentro con el niño).",
     "Calma, asombro, quizá emoción. La primera señal.",
     "Hace la visualización y, al salir, escribe en la libreta lo que llegó (3 palabras, un dibujo o una frase) antes de mirar el celular.",
     "10–12 min", "Pasa de la idea a la experiencia. Es la primera vez que la persona «siente» WAKE UP!.", "FM-01 → 04", "BORRADOR", "", ""),
    ("04", 4, "2 · RECUERDA", "AUDIO CINEMATOGRÁFICO", "CREAR ES SER",
     "Expandir la filosofía: todos somos creadores y crear no es hacer algo, es ser.",
     "Expansión. «Crear es más grande que el arte».",
     "Escucha y, al final, anota 3 cosas que ha creado en su vida sin llamarlas «creativas» (una relación, una solución, un negocio…).",
     "4–6 min", "Rompe el mito de que la creatividad es solo para artistas y amplía el territorio.", "05", "IDEA", "", ""),
    ("05", 5, "2 · RECUERDA", "RAW VIDEO", "¿QUÉ HACÍAS DURANTE HORAS?",
     "Conexión directa: la infancia de Camilo y la pregunta que abre la misión.",
     "Ternura, memoria, ganas de volver a algo.",
     "Recibe la pregunta y la misión 01.",
     "5–8 min", "Hace de puente entre la reflexión y la acción: entrega la misión.", "06", "BORRADOR", "", ""),
    ("06", 6, "2 · RECUERDA", "EJERCICIO CREATIVO", "MISIÓN 01 · EL NIÑO QUE FUISTE",
     "Convertir el aprendizaje en acción, con la estructura de un videojuego: misión, reglas, niveles y recompensa.",
     "Juego, algo de ridículo, liberación.",
     "Lista 10 cosas que hacía de niño, elige una, la juega 30 min sin celular y crea algo a partir de eso.",
     "30–60 min (offline)", "Primera creación real, en el mundo físico. Sin esto, todo se queda en teoría.", "07", "BORRADOR", "", ""),
    ("07", 7, "2 · RECUERDA", "EXPERIENCIA GUIADA", "EL ENCUENTRO",
     "Profundizar emocionalmente: conversar con el niño interior y escuchar qué quiere crear.",
     "Emoción profunda, reconciliación.",
     "Hace la experiencia y después escribe una CARTA A SU NIÑO INTERIOR.",
     "12–15 min", "Es el corazón emocional del programa: aquí se sella el vínculo con el niño.", "08", "IDEA", "", "Versión cinematográfica: más diseño sonoro que la 03."),
    ("08", 8, "3 · JUEGA / CREA", "LIBRO + RAW VIDEO", "CUALQUIERA PUEDE CREAR",
     "Introducir el manifiesto y las páginas del libro, leídas o comentadas por Camilo en RAW.",
     "Convicción, pertenencia.",
     "Lee el manifiesto y empieza el reto de 7 días de ESCRITURA AL DESPERTAR (antes del celular).",
     "8–12 min + reto 7 días", "Crea un hábito que sostiene el despertar después del programa.", "09", "IDEA", "", ""),
    ("09", 9, "3 · JUEGA / CREA", "RAW VIDEO", "UNA IDEA SE AGRADECE HACIÉNDOLA",
     "Integración: pasar de la idea a la acción, aunque sea imperfecta.",
     "Empuje, valentía.",
     "Misión: hacer realidad una idea pequeña en menos de 24 horas y terminarla aunque sea imperfecta.",
     "5–8 min", "Materializar. Cierra la brecha entre inspiración y creación.", "10", "IDEA", "", ""),
    ("10", 10, "3 · JUEGA / CREA", "AUDIO CINEMATOGRÁFICO", "POR AMOR AL ARTE",
     "Crear por amor, no por likes, y empezar a mover a la persona hacia la parte final.",
     "Liviandad, pureza, gratitud.",
     "Escucha y termina con la transición FM hacia el cierre.",
     "4–6 min + FM", "Limpia el porqué: recupera la motivación interna antes del cierre.", "FM-03 → 11", "IDEA", "", ""),
    ("11", 11, "4 · CRECE", "RAW AUDIO / PODCAST", "CONVERSACIÓN",
     "Una conversación profundamente humana (Camilo con otra persona) sobre el miedo, crear por amor y el niño interior.",
     "Compañía. «No estoy solo en esto».",
     "Escucha, idealmente caminando.",
     "20–40 min", "Humaniza y demuestra que el despertar le pasa a cualquiera.", "12", "IDEA", "", "Definir invitado/a."),
    ("12", 12, "4 · CRECE", "EXPERIENCIA GUIADA", "DESPIERTA",
     "Experiencia final de integración: CREER → CREAR → CRECER.",
     "Plenitud, claridad, un nuevo comienzo.",
     "Hace la experiencia y escribe su compromiso creativo, lo primero que va a crear.",
     "12–15 min", "Integra todo el recorrido en el cuerpo y en una decisión.", "FINAL-V", "IDEA", "", ""),
    ("FINAL-V", 13, "4 · CRECE", "VIDEO CINEMATOGRÁFICO", "EL NIÑO ESTÁ DESPIERTO",
     "El símbolo final: el niño está despierto.",
     "Catarsis, belleza.",
     "Mira.",
     "2–4 min", "Es la imagen que se queda en la memoria.", "RAW-FINAL", "IDEA", "", ""),
    ("RAW-FINAL", 14, "EPÍLOGO", "RAW VIDEO", "WAKE UP! · ¿Y AHORA QUÉ?",
     "El despertar ocurrió. ¿Y ahora qué? JUGAR. Puerta a PLAYGROUND.",
     "Ganas de seguir, de jugar.",
     "Recibe la invitación a PLAYGROUND.",
     "2–4 min", "Convierte el final en un comienzo.", "PLAYGROUND", "IDEA", "", ""),
]
ws = sheet("02 MAPA", mapa_h, [11, 8, 16, 24, 30, 42, 34, 42, 14, 38, 13, 14, 26, 26], mapa,
           estado_cols=[12], fmt_col=4)

# ---------------------------------------------------------------- 03 GUIONES
g_h = ["ID", "PIEZA", "#", "BLOQUE", "TIEMPO", "VOZ / TEXTO", "SONIDO / MÚSICA", "VISUAL / IMAGEN", "DIRECCIÓN / NOTAS", "ESTADO"]
G = []


def add(pid, pieza, bloques):
    for i, b in enumerate(bloques, start=1):
        G.append((pid, pieza, i) + tuple(b) + ("BORRADOR",))


add("INTRO", "SEÑAL", [
    ("ESTÁTICA", "0:00–0:15", "(sin voz)", "Estática FM. Un dial busca emisora y pasan fragmentos de voces, canciones y noticias, muy rápidos.", "Negro total.", "Debe sentirse como buscar emisora de noche en un carro viejo."),
    ("LA SEÑAL", "0:15–0:30", "(voz filtrada, como radio lejana) …¿me escuchas?… Si esto te está llegando… es porque todavía hay señal.", "La estática baja y un tono puro se estabiliza.", "Una línea de frecuencia que se mueve.", "La voz tiene que sonar a mensaje encontrado, no a locutor."),
    ("EL PACTO", "0:30–0:55", "(voz limpia, cercana) Esto no es una clase. No vas a tomar apuntes para un examen. Esto es una TRANSMISIÓN. Y va dirigida a una parte de ti que lleva mucho tiempo dormida.", "Entra un pad grave muy suave.", "", "Íntimo, casi al oído."),
    ("PREPARACIÓN", "0:55–1:15", "Ponte audífonos. Busca un lugar donde nadie te interrumpa. Ten cerca una libreta y algo con qué escribir. Los vas a necesitar.", "Pad + una textura de respiración.", "Texto en pantalla: AUDÍFONOS · LIBRETA · SILENCIO", "Da instrucciones sin sonar a manual."),
    ("SINTONÍA", "1:15–1:30", "Estás sintonizando… [PAUSA 2s] WAKE UP.", "Clic de dial, sube el tema principal, golpe y silencio.", "Logo WAKE UP!", "La frecuencia (ej. 88.8) puede ser un símbolo recurrente del universo. Definir."),
])

add("01", "ALGO SE APAGÓ", [
    ("PREGUNTA", "0:00–0:30", "Piensa en la última vez que perdiste la noción del tiempo haciendo algo. [PAUSA 2s] No viendo algo. HACIENDO algo.", "Silencio y luego lluvia suave.", "—", "Arranca sin música. Que la pregunta pese."),
    ("EL RECUERDO", "0:30–1:15", "Quizás fue ayer. Quizás hace años. Quizás tenías siete u ocho años y estabas en el suelo de tu cuarto, construyendo un mundo con lo que tuvieras a mano: fichas, cajas, palos, crayones, una sábana convertida en castillo.", "Ambiente de niños jugando a lo lejos (patio de colegio) y piano suave.", "—", "Voz cálida. Dejar espacio entre cada objeto."),
    ("LO VIVÍAS", "1:15–1:40", "Nadie te había explicado qué era la creatividad. No te hacía falta. La estabas VIVIENDO.", "Piano.", "—", ""),
    ("LO QUE PASÓ", "1:40–2:30", "Después pasó algo. No en un solo día. Poco a poco. Aprendiste que había respuestas correctas. Que había que colorear sin salirse de la línea. Que ciertas cosas «no daban plata». Que era mejor no hacer el ridículo. Que otros lo hacían mejor que tú.", "Los niños se desvanecen y entran la ciudad, el tráfico, notificaciones y un reloj. El piano se desafina levemente.", "—", "Es la parte más dura. Sin dramatizar: con honestidad."),
    ("SE DURMIÓ", "2:30–2:55", "Y esa parte de ti que jugaba, que preguntaba, que inventaba… no murió. [PAUSA 2s] Se quedó dormida.", "Todo se corta de golpe. Silencio.", "—", "El silencio es parte del guion: 3 s mínimo."),
    ("¿QUÉ ES?", "2:55–3:30", "Entonces… ¿qué es la creatividad? La mayoría piensa que es un talento. Algo que tienen los artistas, los músicos, los genios. Algo que se tiene… o no se tiene.", "Entra un pulso bajo, tipo latido.", "—", ""),
    ("TODOS LOS NIÑOS", "3:30–4:00", "Pero mira a cualquier niño de cinco años. Todos dibujan. Todos cantan. Todos inventan historias. Ninguno se pregunta si es «creativo».", "Risa de niño, breve.", "—", "Observación poética, no dato científico (ver 08 INVESTIGACIÓN)."),
    ("DEFINICIÓN", "4:00–4:40", "La creatividad no es un don para unos pocos. Es la capacidad humana de imaginar algo que todavía no existe… y traerlo al mundo. Una canción. Una empresa. Una solución. Una conversación difícil. Una nueva forma de vivir.", "La música crece: cuerdas y sintetizador cálido.", "—", "Con cada ejemplo, un pequeño acento sonoro."),
    ("RECORDAR", "4:40–5:10", "Por eso esto no se trata de aprender algo nuevo. Se trata de RECORDAR algo que ya sabías. [PAUSA 1s] Si algo necesita despertarse… es porque alguna vez estuvo despierto.", "Clímax musical.", "—", ""),
    ("CIERRE", "5:10–5:30", "La creatividad no se aprende. [PAUSA 1s] Se recuerda.", "Corte, estática FM y dial hacia la pieza 02.", "—", "Frase central. Nada después."),
])

add("02", "EL DESPERTAR CREATIVO", [
    ("APERTURA", "0:00–0:45", "Empezar en medio de un pensamiento. Ej.: «No sé muy bien cómo empezar esto, así que voy a empezar por la verdad…»", "Sonido directo, sin música.", "Plano cercano, luz natural, cámara fija o en mano.", "RAW = puntos guía, NO texto para memorizar. Se pueden dejar los errores."),
    ("ANTES", "0:45–2:30", "PREGUNTAS GUÍA: ¿En qué momento sentiste que la creatividad se te había apagado? ¿Qué estabas haciendo en esa época? ¿Qué te decías a ti mismo?", "—", "", "Contar una escena concreta (un lugar, un día), no un resumen."),
    ("EL QUIEBRE", "2:30–4:30", "PREGUNTAS GUÍA: ¿Qué pasó? ¿Cuál fue el momento o la época del despertar? ¿Qué cambió en tu forma de ver la realidad?", "—", "", "Lo espiritual se cuenta como experiencia propia, sin imponérselo a nadie."),
    ("LO QUE DESCUBRÍ", "4:30–6:30", "Llegar a la idea de que crear no es hacer algo, crear es SER. Contar una anécdota concreta de cuando lo sentiste.", "—", "", "Frase ancla: «No haces cosas creativas. Eres creador»."),
    ("POR QUÉ EXISTE WAKE UP!", "6:30–7:30", "«Hice esto porque…». Por qué decidiste compartirlo y qué te gustaría que le pasara a la persona que está viendo.", "—", "", "Sin vender nada."),
    ("INVITACIÓN", "7:30–8:30", "«No te voy a pedir que me creas. Te voy a pedir que lo pruebes. Lo que viene ahora es una experiencia. Busca un lugar cómodo, ponte audífonos y ten tu libreta cerca».", "Al final, un leve zumbido FM.", "Mirada a cámara y corte a negro.", "Prepara directamente la 03."),
])

add("03", "LLAMAR A LO QUE DUERME", [
    ("AVISO", "0:00–0:20", "Esta es una experiencia de relajación y visualización. No la escuches mientras manejas u operas algo. No reemplaza ninguna terapia. Si en algún momento quieres abrir los ojos, puedes hacerlo.", "Silencio.", "Texto en pantalla, el mismo aviso.", "Obligatorio en todas las experiencias guiadas."),
    ("LLEGADA", "0:20–1:30", "Siéntate o acuéstate como estés más cómodo. [PAUSA 3s] Cuando estés listo, cierra los ojos. [PAUSA 3s] Inhala por la nariz… lento… [PAUSA 3s] sostén… [PAUSA 2s] y suelta por la boca. [x3]", "Pad muy suave y respiración guía.", "—", "Voz lenta, un 20% más lenta de lo normal."),
    ("CUERPO", "1:30–3:00", "Siente el peso de tu cuerpo. Los pies… las piernas… la espalda apoyada… los hombros que bajan… la mandíbula que se suelta… la frente. [PAUSA 5s] No tienes que hacer nada. Solo estar aquí.", "Pad.", "—", ""),
    ("EL DIAL", "3:00–4:30", "Imagina frente a ti una radio antigua. Tiene un dial grande con números: son los años de tu vida. [PAUSA 3s] Pon tu mano sobre el dial y empieza a girarlo hacia atrás. [PAUSA 2s] Con cada número que baja, retrocedes un año… y otro… y otro… [PAUSA 5s] Hasta que la aguja se detiene sola en una edad en la que jugabas.", "Clics suaves de dial y estática muy bajita que se va limpiando.", "—", "Motivo de la radio = identidad FM del universo."),
    ("EL LUGAR", "4:30–6:00", "Estás en un lugar donde jugabas de niño. [PAUSA 3s] Mira a tu alrededor. ¿Qué colores hay? [PAUSA 4s] ¿Qué olor? [PAUSA 4s] ¿Qué se escucha? [PAUSA 4s] ¿Qué tienes en las manos? [PAUSA 6s]", "Ambiente muy lejano e impreciso de infancia (pájaros, una casa).", "—", "No describir el lugar: cada persona pone el suyo."),
    ("EL NIÑO", "6:00–7:30", "En algún lugar de ese espacio hay alguien. Eres tú, con esa edad. Está haciendo algo, muy concentrado, sin saber que lo estás mirando. [PAUSA 4s] Obsérvalo. ¿Qué está haciendo? [PAUSA 8s] No le digas nada todavía. Solo recuerda cómo se sentía ser así: sin prisa, sin juicio, sin preguntarse si lo que hace vale la pena. [PAUSA 8s]", "Entra una melodía simple, casi de caja musical.", "—", "Momento emocional. No apurar."),
    ("LA MIRADA", "7:30–8:30", "El niño levanta la mirada y te ve. [PAUSA 3s] ¿Qué ves en su cara? [PAUSA 5s] Hay algo que quiere mostrarte. Déjalo. [PAUSA 10s]", "Melodía.", "—", ""),
    ("LA PREGUNTA", "8:30–9:30", "Antes de irte, pregúntale: ¿qué quieres crear? [PAUSA 3s] No pienses la respuesta. Deja que llegue. Puede ser una imagen, una palabra, un color, un sonido. [PAUSA 20s]", "Casi silencio.", "—", "Silencio largo de verdad. No rellenar."),
    ("LA PROMESA", "9:30–10:00", "Dile que vas a volver. Que esta vez no lo vas a dejar dormido tanto tiempo. [PAUSA 5s]", "La melodía se resuelve.", "—", ""),
    ("REGRESO", "10:00–11:00", "Pon tu mano otra vez en el dial y gíralo hacia adelante, año por año, hasta el día de hoy. [PAUSA 4s] Voy a contar del 1 al 5 y con cada número vuelves un poco más. 1… 2, siente tu cuerpo… 3, mueve los dedos… 4, respira profundo… 5, abre los ojos cuando quieras.", "Clics de dial hacia adelante y estática que sube y se limpia.", "—", ""),
    ("PRIMERA SEÑAL", "11:00–11:40", "Antes de hacer cualquier otra cosa, antes de mirar el celular, abre tu libreta y escribe lo que llegó. Tres palabras, un dibujo, una frase. No lo juzgues. Es tu primera SEÑAL.", "Tema WAKE UP! muy suave.", "Texto: ESCRIBE LO QUE LLEGÓ", "Es la acción real de la pieza. Esa libreta acompaña todo el programa."),
])

add("05", "¿QUÉ HACÍAS DURANTE HORAS?", [
    ("APERTURA", "0:00–1:00", "Contar, de forma concreta, algo que Camilo hacía durante horas de niño. Si hay un objeto o una foto de esa época, mostrarlo.", "Directo.", "Plano cercano. Objeto o foto real en cámara.", "RAW. Puntos guía."),
    ("QUÉ SENTÍA", "1:00–2:30", "¿Qué sentías mientras lo hacías? ¿Se te pasaba el tiempo? ¿Alguien te decía que lo dejaras?", "", "", "Nombrar el flow sin tecnicismos: «perder la noción del tiempo»."),
    ("CUÁNDO LO DEJÉ", "2:30–4:00", "¿Cuándo y por qué lo dejaste? ¿Qué lo reemplazó?", "", "", "Conecta con la 01 (lo que pasó)."),
    ("LA PREGUNTA", "4:00–5:00", "Mirando a cámara: «Ahora te pregunto a ti: ¿qué hacías durante horas cuando eras niño?»", "", "", "Pausa real después de la pregunta."),
    ("ENTREGA DE MISIÓN", "5:00–6:00", "«Lo que viene no es para verlo. Es para hacerlo. Te voy a dar tu primera misión».", "Sonido de «misión desbloqueada».", "Transición gráfica a la pantalla de misión (06).", ""),
])

add("06", "MISIÓN 01 · EL NIÑO QUE FUISTE", [
    ("PANTALLA DE MISIÓN", "—", "MISIÓN 01: EL NIÑO QUE FUISTE. Dificultad: ★☆☆. Duración: 30–60 min. Recompensa: desbloquear la pieza 07.", "Sonido de juego (arcade suave o 8 bits).", "Interfaz tipo videojuego con la estética WAKE UP!.", "Puede ser video corto, página interactiva o PDF."),
    ("OBJETIVO", "—", "Volver a jugar, durante 30 minutos, a algo que hacías de niño y crear algo a partir de eso.", "", "", ""),
    ("REGLAS", "—", "1. No tiene que quedar bien. 2. No se puede monetizar ni subir para likes. 3. Celular en modo avión (solo para la foto final). 4. Si te da vergüenza, vas bien.", "", "", "Las reglas son parte del juego y le quitan presión."),
    ("EQUIPO", "—", "Tu libreta, algo con qué escribir y lo que necesites para tu juego (colores, legos, una pelota, papel, tu voz…).", "", "", ""),
    ("NIVEL 1 · RECORDAR", "5 min", "Escribe en tu libreta 10 cosas que hacías durante horas cuando eras niño. No filtres.", "", "", ""),
    ("NIVEL 2 · JUGAR", "30 min", "Elige una y hazla hoy durante 30 minutos. Versión adulta permitida, pero con reglas de niño: sin objetivo, sin juicio.", "", "", ""),
    ("NIVEL 3 · CREAR", "10–20 min", "A partir de lo que jugaste, crea UNA cosa: un dibujo, una frase, una canción de 20 segundos, un objeto, una historia. Termínala.", "", "", "Jefe final: terminar algo imperfecto."),
    ("REGISTRO", "5 min", "Tómale una foto a lo que creaste y escribe una línea: «Hoy recordé que…».", "", "", "Opcional: galería de la comunidad a futuro (PLAYGROUND)."),
    ("RECOMPENSA", "—", "MISIÓN COMPLETA. Desbloqueaste: 07 · EL ENCUENTRO.", "Sonido de victoria.", "", ""),
])

sheet("03 GUIONES", g_h, [10, 26, 5, 20, 11, 70, 36, 30, 40, 13], G, estado_cols=[10])

# ---------------------------------------------------------------- 04 EJERCICIOS
e_h = ["ID", "NOMBRE", "TIPO", "MISIÓN (1 línea)", "INSTRUCCIONES", "TIEMPO", "MATERIALES", "RESULTADO / QUÉ SE LLEVA", "DIFICULTAD", "PIEZA", "ESTADO"]
ej = [
    ("EJ-01", "Primera señal", "Escritura", "Escribe lo que llegó antes de mirar el celular.", "Al terminar la experiencia 03, escribe 3 palabras, un dibujo o una frase. Sin juzgar.", "2 min", "Libreta", "Primer registro en la libreta WAKE UP!", "★☆☆", "03", "BORRADOR"),
    ("EJ-02", "El niño que fuiste", "Juego / Acción", "Vuelve a jugar 30 min a algo de tu infancia y crea algo a partir de eso.", "Ver el guion de la 06 (niveles 1–3).", "30–60 min", "Libreta + materiales del juego", "Una creación terminada + foto", "★☆☆", "06", "BORRADOR"),
    ("EJ-03", "Carta a tu niño interior", "Escritura", "Escríbele a quien fuiste.", "Después de la 07, escribe una carta a tu niño interior: qué le pides perdón, qué le prometes y qué quieres crear con él.", "15–20 min", "Libreta", "Una carta", "★★☆", "07", "IDEA"),
    ("EJ-04", "Carta a tu yo bloqueado", "Escritura", "Háblale a la parte de ti que tiene miedo.", "Escribe una carta al yo que no se atreve: qué teme, qué le dirías, qué necesita escuchar.", "15 min", "Libreta", "Una carta", "★★☆", "Por definir", "IDEA"),
    ("EJ-05", "Escritura al despertar (7 días)", "Hábito / Escritura", "Siete mañanas escribiendo antes del celular.", "Al despertar, con la libreta al lado de la cama, escribe sin parar entre 1 y 3 páginas antes de mirar el celular. Todo vale.", "10–20 min / día", "Libreta en la mesa de noche", "Hábito + descarga de ideas", "★★☆", "08", "IDEA"),
    ("EJ-06", "Diario de sueños", "Escritura", "Atrapa los sueños antes de que se vayan.", "Anota lo que recuerdes de tus sueños apenas te despiertes: imágenes, colores, frases.", "5 min / día", "Libreta", "Banco de imágenes", "★☆☆", "08", "IDEA"),
    ("EJ-07", "Descarga sin juicio", "Escritura", "10 minutos de ideas sin filtro.", "Temporizador de 10 minutos para escribir todas las ideas que lleguen, incluso las absurdas. No tachar.", "10 min", "Libreta + temporizador", "Lista de ideas crudas", "★☆☆", "Por definir", "IDEA"),
    ("EJ-08", "Crea algo cada día", "Hábito", "Una creación diaria, por pequeña que sea.", "Durante 7 días crea algo pequeño cada día: una frase, una foto pensada, un dibujo, una receta.", "5–30 min / día", "Libre", "7 creaciones", "★★☆", "Por definir", "IDEA"),
    ("EJ-09", "Cita contigo", "Exploración", "Sal solo a un lugar que te dé curiosidad.", "Una vez por semana, 1–2 horas solo en un lugar nuevo (museo, mercado, parque, ferretería…). Sin agenda.", "1–2 h", "Nada", "Estímulos + ideas", "★★☆", "Por definir", "IDEA"),
    ("EJ-10", "Primera vez", "Observación", "Mira algo cotidiano como si nunca lo hubieras visto.", "Elige un objeto de tu casa y descríbelo durante 5 minutos como si fueras un extraterrestre que lo ve por primera vez.", "10 min", "Libreta", "Mirada renovada", "★☆☆", "04 / 05", "IDEA"),
    ("EJ-11", "Crear con límites", "Juego", "Crea algo con solo 3 elementos.", "Elige 3 materiales, colores, palabras o sonidos y crea algo usando solo eso.", "20 min", "3 elementos a elección", "Una creación", "★★☆", "Por definir", "IDEA"),
    ("EJ-12", "Termina lo imperfecto", "Acción", "Termina algo que dejaste a medias.", "Busca una idea o proyecto abandonado y termínalo HOY, aunque quede imperfecto.", "30–60 min", "El proyecto", "Algo TERMINADO", "★★★", "09", "IDEA"),
    ("EJ-13", "Idea en 24 horas", "Acción", "Haz realidad una idea pequeña en menos de un día.", "Elige una idea que quepa en 24 horas y hazla. «Una idea se agradece haciéndola».", "< 24 h", "Libre", "Una idea materializada", "★★★", "09", "IDEA"),
    ("EJ-14", "Mapa del flow", "Observación", "Descubre cuándo entras en flow.", "Durante una semana, anota cada vez que pierdas la noción del tiempo: qué hacías, dónde, a qué hora y con quién.", "1 semana", "Libreta", "Tu mapa personal de flow", "★★☆", "Por definir", "IDEA"),
    ("EJ-15", "Solo por curiosidad", "Exploración", "Aprende algo que no sirva para nada.", "Explora durante 1 hora un tema solo porque te da curiosidad, sin utilidad y sin objetivo.", "1 h", "Libre", "Curiosidad reactivada", "★☆☆", "10", "IDEA"),
    ("EJ-16", "Compromiso creativo", "Escritura / Acción", "Escribe lo primero que vas a crear.", "Después de la 12, escribe qué vas a crear, cuándo empiezas y cuál es el primer paso pequeño.", "10 min", "Libreta", "Un compromiso concreto", "★★☆", "12", "IDEA"),
]
sheet("04 EJERCICIOS", e_h, [8, 26, 16, 40, 60, 14, 24, 28, 11, 12, 13], ej, estado_cols=[11])

# ---------------------------------------------------------------- 05 FM
fm_h = ["ID", "UBICACIÓN", "FRECUENCIA", "MENSAJE / TEXTO", "DISEÑO SONORO", "DURACIÓN", "FUNCIÓN", "ESTADO"]
fm = [
    ("FM-00", "Inicio (INTRO)", "88.8", "…¿me escuchas? Si esto te está llegando, es porque todavía hay señal.", "Dial buscando, voces cruzadas, tono que se estabiliza.", "incluido en INTRO", "Abrir el universo.", "BORRADOR"),
    ("FM-01", "03 → 04", "Por definir", "(interferencia) …señal recibida… el receptor está despertando…", "Estática corta, pitido de confirmación.", "10–15 s", "Confirmar que algo empezó a moverse después de la primera experiencia.", "IDEA"),
    ("FM-02", "07 → 08", "Por definir", "Fragmento de una voz infantil (¿la de Camilo de niño?) diciendo una frase simple.", "Cambio de canal y cinta vieja.", "10–20 s", "Hacer eco del encuentro con el niño.", "IDEA"),
    ("FM-03", "10 → 11", "Por definir", "Transmisión final: «Última frecuencia antes del despertar».", "Dial que sube de frecuencia y la señal cada vez más limpia.", "20–30 s", "Mover al usuario hacia la parte final.", "IDEA"),
    ("FM-04", "12 → FINAL-V", "Por definir", "Silencio total y luego la señal perfectamente limpia.", "Ausencia de estática por primera vez.", "5–10 s", "Simbolizar que ya no hay interferencia: el niño despertó.", "IDEA"),
]
sheet("05 TRANSICIONES FM", fm_h, [8, 18, 12, 50, 40, 16, 44, 13], fm, estado_cols=[8])

# ---------------------------------------------------------------- 06 PRODUCCIÓN
p_h = ["ID", "PIEZA", "FORMATO", "GUION", "GRABACIÓN", "EDICIÓN", "MÚSICA / SFX", "MEZCLA / COLOR", "REVISIÓN", "FINAL", "FECHA OBJETIVO", "RESPONSABLE", "LINK ARCHIVO", "NOTAS"]
prod = []
for r in mapa:
    guion = "EN PROCESO" if r[11] == "BORRADOR" else "PENDIENTE"
    prod.append((r[0], r[4], r[3], guion, "PENDIENTE", "PENDIENTE", "PENDIENTE", "PENDIENTE", "PENDIENTE", "PENDIENTE", "", "Camilo", "", ""))
sheet("06 PRODUCCIÓN", p_h, [11, 30, 24, 13, 13, 13, 13, 14, 13, 13, 15, 14, 26, 30], prod,
      prod_cols=list(range(4, 11)), fmt_col=3)

# ---------------------------------------------------------------- 07 FRASES
f_h = ["FRASE", "USO", "PIEZA", "ORIGEN"]
frases = [
    ("LA CREATIVIDAD NO SE APRENDE. SE RECUERDA.", "Frase central / cierre", "01, global", "Notas de Camilo"),
    ("YOU DON'T LEARN CREATIVITY. YOU REMEMBER IT.", "Frase central (EN)", "Global", "Notas de Camilo"),
    ("DESPIERTA. RECUERDA. JUEGA. CREA.", "Tagline", "Global", "Contexto maestro"),
    ("Crear no es hacer algo… crear es ser.", "Principio maestro", "02, 04", "Notas de Camilo"),
    ("No haces cosas creativas. Eres creador.", "Copy / RAW", "02, 04", "Contexto maestro"),
    ("TODOS SOMOS CREADORES.", "Título / manifiesto", "04, 08", "Contexto maestro"),
    ("Si algo necesita despertarse, es porque alguna vez estuvo despierto.", "Guion", "01", "Contexto maestro"),
    ("Una idea se agradece haciéndola realidad.", "Guion / misión", "09", "Notas de Camilo"),
    ("CREER → CREAR → CRECER", "Estructura / cierre", "12", "Notas de Camilo"),
    ("POR AMOR AL ARTE", "Título", "10", "Notas de Camilo"),
    ("EL NIÑO ESTÁ DESPIERTO.", "Símbolo final", "FINAL-V", "Contexto maestro"),
    ("¿Y AHORA QUÉ? JUGAR.", "Puente a PLAYGROUND", "RAW-FINAL", "Contexto maestro"),
    ("Si esto te está llegando, es porque todavía hay señal.", "Guion", "INTRO", "Propuesta"),
    ("Si te da vergüenza, vas bien.", "Regla de misión", "06", "Propuesta"),
    ("No te voy a pedir que me creas. Te voy a pedir que lo pruebes.", "RAW", "02", "Propuesta"),
]
sheet("07 FRASES", f_h, [60, 26, 14, 20], frases)

# ---------------------------------------------------------------- 08 INVESTIGACIÓN
i_h = ["TEMA / AFIRMACIÓN", "TIPO", "QUÉ DICE LA EVIDENCIA (resumen)", "FUENTE A REVISAR", "¿SE PUEDE CITAR COMO DATO?", "PIEZA"]
inv = [
    ("«98% de los niños de 5 años son genios creativos y solo el 2% de los adultos» (prueba de la NASA)", "POPULAR · VERIFICAR",
     "Viene de George Land (libro «Breakpoint and Beyond», 1992). Se cita muchísimo, pero no se publicó como estudio revisado por pares y su metodología es poco clara.",
     "Land & Jarman (1992), Breakpoint and Beyond", "NO como dato duro. Sí como «se dice que…» o como metáfora.", "01"),
    ("Los puntajes de pensamiento creativo han bajado en las últimas décadas", "EVIDENCIA · VERIFICAR",
     "Kyung Hee Kim analizó los puntajes del test de Torrance (TTCT) y encontró una caída desde los años 90 en EE. UU.",
     "Kim, K. H. (2011), Creativity Research Journal, «The Creativity Crisis»", "Sí, con fuente y matiz (una población, un test).", "01"),
    ("Flow / perder la noción del tiempo", "EVIDENCIA",
     "Es un constructo psicológico muy estudiado (Csikszentmihalyi): reto equilibrado con habilidad, concentración total y alteración de la percepción del tiempo.",
     "Csikszentmihalyi, M. (1990), Flow", "Sí.", "05, EJ-14"),
    ("Incubación: las ideas aparecen cuando dejas de pensar en el problema", "EVIDENCIA",
     "Un metaanálisis encuentra un efecto de incubación positivo, sobre todo en tareas de pensamiento divergente.",
     "Sio & Ormerod (2009), Psychological Bulletin", "Sí, con matiz.", "04, 10"),
    ("Mente divagante / red neuronal por defecto y creatividad", "EVIDENCIA · MATIZ",
     "Los estudios de neuroimagen asocian el pensamiento creativo a la interacción entre la red por defecto y las redes de control ejecutivo, no a la red por defecto sola.",
     "Beaty et al. (2018), PNAS", "Sí, sin simplificar a «el cerebro creativo».", "04"),
    ("El estado entre la vigilia y el sueño (hipnagogia) favorece las ideas", "EVIDENCIA",
     "Pasar un rato en la fase N1 del sueño aumentó la probabilidad de descubrir una regla oculta (inspirado en el método de Edison).",
     "Lacaux et al. (2021), Science Advances", "Sí.", "08, EJ-05, EJ-06"),
    ("Escritura al despertar («páginas matutinas»)", "PRÁCTICA",
     "Es una práctica popularizada por Julia Cameron. Hay evidencia general sobre la escritura expresiva, pero poca investigación directa sobre las páginas matutinas.",
     "Cameron (1992), The Artist's Way; Pennebaker (escritura expresiva)", "Como práctica recomendada, no como «está probado que…».", "08"),
    ("Juego simulado en la infancia y creatividad", "EVIDENCIA · MATIZ",
     "Hay estudios que asocian el juego de simulación con el pensamiento divergente, pero las revisiones advierten que la evidencia causal es débil.",
     "Russ (2014); Lillard et al. (2013), Psychological Bulletin", "Con cuidado: «se asocia», no «causa».", "05, 06"),
    ("La creatividad se recuerda / el despertar / crear es ser", "FILOSOFÍA · METÁFORA",
     "Es el marco narrativo y filosófico de WAKE UP!. No hace falta probarlo: se presenta como visión, no como dato.",
     "—", "No aplica (es la voz del proyecto).", "Global"),
    ("Hipnosis / experiencias guiadas", "PRECAUCIÓN",
     "Se presentan como experiencias de relajación y visualización, nunca como tratamiento médico o psicológico. Llevan aviso de seguridad.",
     "—", "No hacer promesas terapéuticas.", "03, 07, 12"),
]
sheet("08 INVESTIGACIÓN", i_h, [48, 20, 60, 40, 32, 14], inv)

# ---------------------------------------------------------------- 09 IDEAS
id_h = ["FECHA", "IDEA", "RELACIONADO CON", "PRIORIDAD", "ESTADO"]
ideas = [
    ("", "PLAYGROUND: la siguiente fase, un espacio para jugar, explorar y crear después del despertar.", "RAW-FINAL", "ALTA", "IDEA"),
    ("", "Libreta física WAKE UP! como objeto del programa (¿kit con la compra?).", "03, 08", "MEDIA", "IDEA"),
    ("", "La frecuencia 88.8 (u otra) como símbolo recurrente del universo.", "FM", "MEDIA", "IDEA"),
    ("", "Usar un audio real de Camilo de niño en la FM-02.", "FM-02, 07", "MEDIA", "IDEA"),
    ("", "Galería de la comunidad con las creaciones de la Misión 01.", "06, PLAYGROUND", "BAJA", "IDEA"),
    ("", "El libro «Cualquiera puede crear»: índice y primeras páginas.", "08", "ALTA", "IDEA"),
    ("", "Alinear la web (wakeup.html / módulos) con esta nueva estructura de 12 piezas + intro + final.", "Web", "ALTA", "IDEA"),
]
ws = sheet("09 IDEAS", id_h, [12, 80, 20, 12, 13], ideas, estado_cols=[5])
dv = DataValidation(type="list", formula1='"ALTA,MEDIA,BAJA"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("D2:D300")

wb.save(OUT)
print("OK", OUT)
